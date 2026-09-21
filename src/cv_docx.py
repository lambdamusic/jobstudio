"""
cv_docx.py — ATS-friendly Word (.docx) rendering for CV data.

Consumes the same parsed dicts that render.py produces (parse_functional /
parse_variant) and writes a deliberately plain, single-column .docx:
real heading styles, standard bullet lists, real hyperlinks, no columns,
no text boxes, no pseudo-element glyphs. This is the artifact to upload to
applicant tracking systems; the Chrome-rendered PDF stays the "designed"
version for human readers.
"""

import html
import re

import identity

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x1B, 0x29, 0x48)
GOLD = RGBColor(0xC4, 0x9B, 0x2E)
BODY = RGBColor(0x1A, 0x1A, 0x1A)
NAVY_HEX = "1B2948"
GOLD_HEX = "C49B2E"

# Inline HTML (as emitted by render._inline) → docx runs.
_RUN_RE = re.compile(
    r'<a href="([^"]*)">(.*?)</a>'
    r'|<strong>(.*?)</strong>'
    r'|<em>(.*?)</em>'
    r'|([^<]+)',
    re.DOTALL,
)


def _plain(s: str) -> str:
    """Strip any HTML tags and unescape entities → plain text."""
    return html.unescape(re.sub(r'<[^>]+>', '', s or ''))


def _add_hyperlink(paragraph, url: str, text: str, bold=False, italic=False):
    # Word's default template has no "Hyperlink" style, so referencing it via
    # w:rStyle renders as plain, uncoloured text — invisible as a link. Set the
    # navy colour + underline directly instead.
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), NAVY_HEX)
    rpr.append(color)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rpr.append(u)
    if bold:
        rpr.append(OxmlElement('w:b'))
    if italic:
        rpr.append(OxmlElement('w:i'))
    run.append(rpr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _bottom_border(paragraph, color_hex: str, sz: str = '6', space: str = '4'):
    """Add a single bottom border/rule to *paragraph* (used for section rules)."""
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), sz)
    bottom.set(qn('w:space'), space)
    bottom.set(qn('w:color'), color_hex)
    pbdr.append(bottom)
    paragraph._p.get_or_add_pPr().append(pbdr)


def _add_runs(paragraph, s: str, bold=False, italic=False, color=None):
    """Add runs to *paragraph* from an inline-HTML fragment.

    *color*, if given, applies to plain-text and italic runs (bold/link runs
    already default to navy — see below).
    """
    for m in _RUN_RE.finditer(s or ''):
        if m.group(1) is not None:
            _add_hyperlink(paragraph, html.unescape(m.group(1)),
                           html.unescape(m.group(2)), bold=bold, italic=italic)
        elif m.group(3) is not None:
            # **bold** is only ever used for label/employer prefixes in this CV's
            # markdown convention, so colour it navy for a consistent accent.
            r = paragraph.add_run(html.unescape(m.group(3)))
            r.bold = True
            r.italic = italic
            r.font.color.rgb = color or NAVY
        elif m.group(4) is not None:
            r = paragraph.add_run(html.unescape(m.group(4)))
            r.italic = True
            r.bold = bold
            if color:
                r.font.color.rgb = color
        elif m.group(5) is not None:
            r = paragraph.add_run(html.unescape(m.group(5)))
            r.bold = bold
            r.italic = italic
            if color:
                r.font.color.rgb = color


# ── Document scaffolding ────────────────────────────────────────────────────

def _new_doc(data: dict) -> Document:
    doc = Document()

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.06

    h2 = doc.styles['Heading 2']
    h2.font.name = 'Calibri'
    h2.font.size = Pt(11.5)
    h2.font.bold = True
    h2.font.color.rgb = GOLD
    h2.font.all_caps = True
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True
    # subtle letter-spacing to echo the gold letter-spaced section titles in the
    # HTML/PDF templates (0.75pt, in twentieths-of-a-point per the OOXML unit)
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:val'), '15')
    h2.element.get_or_add_rPr().append(spacing)

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    doc.core_properties.author = _plain(f"{data['first_name']} {data['last_name']}")
    doc.core_properties.title = _plain(data['title'])

    _header(doc, data)
    return doc


def _header(doc: Document, data: dict):
    name_p = doc.add_paragraph()
    name_p.paragraph_format.space_after = Pt(2)
    run = name_p.add_run(_plain(f"{data['first_name']} {data['last_name']}"))
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = NAVY

    if data.get('current_role'):
        role_p = doc.add_paragraph()
        role_p.paragraph_format.space_after = Pt(4)
        r = role_p.add_run(_plain(data['current_role']))
        r.italic = True
        r.bold = True
        r.font.color.rgb = GOLD

    contact_p = doc.add_paragraph()
    contact_p.paragraph_format.space_after = Pt(10)
    for i, item in enumerate(data['contact']):
        if i:
            run = contact_p.add_run('   ·   ')
            run.font.color.rgb = GOLD
            run.bold = True
        if item['url']:
            _add_hyperlink(contact_p, item['url'], item['label'])
        else:
            contact_p.add_run(item['text']).font.color.rgb = NAVY
    _bottom_border(contact_p, NAVY_HEX, sz='18', space='8')


def _section(doc: Document, title: str):
    h = doc.add_heading(title, level=2)
    _bottom_border(h, GOLD_HEX, sz='6', space='4')
    return h


def _summary(doc: Document, data: dict):
    if not data['summary']:
        return
    _section(doc, 'Summary')
    for para in data['summary']:
        _add_runs(doc.add_paragraph(), para)


def _education_skills_pubs(doc: Document, data: dict):
    if data['education']:
        _section(doc, 'Education')
        for e in data['education']:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            deg = p.add_run(_plain(e['degree']))
            deg.bold = True
            deg.font.color.rgb = NAVY
            if e['extra']:
                p.add_run(' ')
                _add_runs(p, e['extra'])
            bits = []
            if e['inst']:
                bits.append(_plain(e['inst']))
            if e['years']:
                bits.append(_plain(e['years']))
            if e['focus']:
                bits.append('Focus: ' + _plain(e['focus']))
            if bits:
                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(5)
                sp.add_run(' · '.join(bits)).italic = True

    if data['skills']:
        _section(doc, 'Skills')
        for s in data['skills']:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            lbl = p.add_run(_plain(s['label']) + ': ')
            lbl.bold = True
            lbl.font.color.rgb = NAVY
            _add_runs(p, s['text'])

    if data['languages']:
        _section(doc, 'Languages')
        doc.add_paragraph(
            ', '.join(f"{l['lang']} ({l['level']})" for l in data['languages'])
        )

    if data['publications']:
        _section(doc, 'Publications')
        for pub in data['publications']:
            _add_runs(doc.add_paragraph(style='List Bullet'), pub, italic=True)


# ── Public builders ────────────────────────────────────────────────────────

def build_functional(data: dict, out_path):
    doc = _new_doc(data)
    _summary(doc, data)

    _section(doc, 'Core Competencies')
    for c in data['competencies']:
        tp = doc.add_paragraph()
        tp.paragraph_format.space_before = Pt(7)
        tp.paragraph_format.space_after = Pt(1)
        tp.paragraph_format.keep_with_next = True
        tr = tp.add_run(_plain(c['title']))
        tr.bold = True
        tr.font.size = Pt(11)
        tr.font.color.rgb = NAVY
        if c['statement']:
            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(2)
            sp.paragraph_format.keep_with_next = True
            _add_runs(sp, c['statement'], italic=True)
        for b in c['bullets']:
            _add_runs(doc.add_paragraph(style='List Bullet'), b)

    _section(doc, 'Experience')
    for e in data['career']:
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(6)
        hp.paragraph_format.space_after = Pt(1)
        hp.paragraph_format.keep_with_next = True
        _add_runs(hp, e['org_html'], bold=True, color=NAVY)
        meta = _plain(e['dates'])
        if e['context']:
            ctx = _plain(e['context'])
            meta = f"{meta} · {ctx}" if meta else ctx
        if meta:
            hp.add_run(f"   ({meta})" if e['dates'] else f"   {meta}").italic = True
        if e.get('summary'):
            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(2)
            sp.paragraph_format.keep_with_next = True
            _add_runs(sp, e['summary'], italic=True)
        for r in e['roles']:
            rp = doc.add_paragraph(style='List Bullet')
            _add_runs(rp, r['role'])
            if r['dates']:
                rp.add_run(f"   ({_plain(r['dates'])})")

    _education_skills_pubs(doc, data)
    doc.save(str(out_path))
    return out_path


_LETTERHEAD_GREY = RGBColor(0x55, 0x55, 0x55)


def build_cover_letter(blocks, out_path):
    """Render parsed cover-letter blocks to a plain business-letter .docx.

    *blocks* is the list produced by render._parse_cover_letter:
    ('title', text) | ('head', frags) | ('hr',) | ('p', frags),
    where frags is [(inline_html, hard_break), ...]. 'title' is a file label and is
    written to the document properties only, not the page. 'head' blocks are the
    letterhead (name, address, email, date) — left-aligned, small, grey.
    """
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Georgia'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.35
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.1)
        section.right_margin = Inches(1.1)
    # .docx metadata travels with the file to the employer — a hardcoded author here
    # put the toolkit author's name inside every user's documents.
    doc.core_properties.author = identity.full_name()
    doc.core_properties.title = 'Cover Letter'

    def _fill(p, frags, grey=False):
        for i, (frag, hard) in enumerate(frags):
            runs_before = len(p.runs)
            _add_runs(p, frag)
            if grey:
                for r in p.runs[runs_before:]:
                    r.font.color.rgb = _LETTERHEAD_GREY
                    r.font.size = Pt(9.5)
            if i < len(frags) - 1:
                if hard:
                    p.add_run().add_break()
                else:
                    p.add_run(' ')

    for block in blocks:
        if block[0] == 'title':
            doc.core_properties.title = block[1]
        elif block[0] == 'head':
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            _fill(p, block[1], grey=True)
        elif block[0] == 'hr':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(12)
            _bottom_border(p, GOLD_HEX, sz='6', space='1')
        else:
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _fill(p, block[1])

    doc.save(str(out_path))
    return out_path


def build_chronological(data: dict, out_path):
    doc = _new_doc(data)
    _summary(doc, data)

    _section(doc, 'Experience')
    for j in data['jobs']:
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(6)
        hp.paragraph_format.space_after = Pt(1)
        hp.paragraph_format.keep_with_next = True
        rr = hp.add_run(_plain(j['role']) + ' — ')
        rr.bold = True
        rr.font.color.rgb = NAVY
        _add_runs(hp, j['company_html'], bold=True, color=NAVY)
        if j['meta']:
            mp = doc.add_paragraph()
            mp.paragraph_format.space_after = Pt(2)
            mp.paragraph_format.keep_with_next = True
            mr = mp.add_run(_plain(j['meta']))
            mr.italic = True
            mr.font.size = Pt(9.5)
        for b in j['bullets']:
            _add_runs(doc.add_paragraph(style='List Bullet'), b)
        if j['tech']:
            _add_runs(doc.add_paragraph(), j['tech'])

    ec = data.get('early_career')
    if ec:
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(6)
        hp.paragraph_format.keep_with_next = True
        r = hp.add_run(_plain(ec['role']))
        r.bold = True
        r.font.color.rgb = NAVY
        if ec.get('meta'):
            mp = doc.add_paragraph()
            mp.paragraph_format.space_after = Pt(2)
            mp.add_run(_plain(ec['meta'])).italic = True
        if ec.get('bullets'):
            for b in ec['bullets']:
                _add_runs(doc.add_paragraph(style='List Bullet'), b)
        elif ec.get('prose'):
            _add_runs(doc.add_paragraph(), ec['prose'])

    _education_skills_pubs(doc, data)
    doc.save(str(out_path))
    return out_path
